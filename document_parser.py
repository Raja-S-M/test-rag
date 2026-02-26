import fitz  # PyMuPDF
import docx
import pptx
import pandas as pd
import io

def parse_pdf(file_bytes: bytes) -> str:
    text = ""
    with fitz.open(stream=file_bytes, filetype="pdf") as doc:
        for page in doc:
            text += page.get_text() + "\n"
    return text

def parse_docx(file_bytes: bytes) -> str:
    doc = docx.Document(io.BytesIO(file_bytes))
    return "\n".join([para.text for para in doc.paragraphs])

def parse_pptx(file_bytes: bytes) -> str:
    prs = pptx.Presentation(io.BytesIO(file_bytes))
    text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
    return text

def parse_excel(file_bytes: bytes) -> str:
    df = pd.read_excel(io.BytesIO(file_bytes))
    return df.to_string()

def parse_csv(file_bytes: bytes) -> str:
    df = pd.read_csv(io.BytesIO(file_bytes))
    return df.to_string()

def parse_document(file_name: str, file_bytes: bytes) -> str:
    ext = file_name.split(".")[-1].lower()
    if ext == "pdf":
        return parse_pdf(file_bytes)
    elif ext in ["doc", "docx"]:
        return parse_docx(file_bytes)
    elif ext == "pptx":
        return parse_pptx(file_bytes)
    elif ext in ["xls", "xlsx"]:
        return parse_excel(file_bytes)
    elif ext == "csv":
        return parse_csv(file_bytes)
    else:
        # Fallback to plain text decoding
        return file_bytes.decode("utf-8", errors="ignore")
